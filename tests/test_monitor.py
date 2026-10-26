import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import requests
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import Workbook

import monitor
import tender_pipeline as pipeline


ENV = {
    "GEMINI_API_KEY": "test-not-a-real-key",
    "GEMINI_MODEL": "test-model",
    "EMAIL_TO": "recipient@example.test",
    "EMAIL_FROM": "sender@example.test",
    "EMAIL_PASSWORD": "test-password",
}
ANALYSIS = {
    "subject": None, "delivery_deadline": None, "warranty": None, "confidential": None,
}
TENDER = {
    "id": "123", "reg_num": "NAT123", "buyer": "Buyer",
    "link": pipeline.PORTAL + "?go=123",
}
DOCS = """<h2>დოკუმენტაცია</h2>
<a href="library/files.php?mode=app&amp;file=42&amp;code=100">spec.pdf</a>"""


def response(payload=None, status=200, text=""):
    result = MagicMock()
    result.status_code = status
    result.text = text
    result.json.return_value = payload
    result.__enter__.return_value = result
    if status >= 400:
        result.raise_for_status.side_effect = requests.HTTPError(f"HTTP {status}")
    return result


def generation(analysis=None, finish="STOP"):
    return response({
        "usageMetadata": {
            "promptTokenCount": 50, "candidatesTokenCount": 20, "totalTokenCount": 80,
        },
        "candidates": [{
            "finishReason": finish,
            "content": {"parts": [{"text": json.dumps(analysis or ANALYSIS)}]},
        }],
    })


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(os.environ, ENV, clear=True)
        self.environment.start()
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.addCleanup(self.environment.stop)
        self.path = Path(self.directory.name) / "state.json"
        os.environ["STATE_PATH"] = str(self.path)
        self.state = pipeline.load_state(self.path)
        self.ai = pipeline.Gemini(self.state, self.path)

    def test_state_roundtrip_and_corruption(self):
        pipeline.save_state(self.path, self.state)
        self.assertEqual(self.state, pipeline.load_state(self.path))
        self.path.write_text('{"version": 2}', encoding="utf-8")
        with self.assertRaises(ValueError):
            pipeline.load_state(self.path)
        self.assertFalse(self.path.with_suffix(".tmp").exists())

    def test_document_links_resolve_against_public_root_and_support_que(self):
        docs = pipeline.discover_documents(DOCS + """
        <a href="library/files.php?mode=que&amp;file=43">sheet.xlsx</a>
        <a href="https://evil.example/files.php?mode=app&amp;file=44">unsafe</a>
        <a href="library/files.php?mode=app&amp;file=42&amp;code=200">duplicate</a>
        """)
        self.assertEqual([d["id"] for d in docs], ["app/42", "que/43"])
        self.assertIn("/public/library/files.php?", docs[0]["url"])
        self.assertNotIn("/library/library/", docs[0]["url"])

    def test_download_rejects_html_and_redirects_and_large_files(self):
        session = MagicMock()
        reply = response()
        session.get.return_value = reply
        reply.iter_content.return_value = [b"<html>session error</html>"]
        with self.assertRaisesRegex(pipeline.ProcessingError, "HTML"):
            pipeline.download_document(session, pipeline.discover_documents(DOCS)[0])
        reply.status_code = 302
        with self.assertRaisesRegex(pipeline.ProcessingError, "redirected"):
            pipeline.download_document(session, pipeline.discover_documents(DOCS)[0])
        reply.status_code = 200
        reply.iter_content.return_value = [b"%PDF-" + b"a" * 20]
        with patch.object(pipeline, "MAX_BYTES", 10):
            with self.assertRaisesRegex(pipeline.ProcessingError, "safety limit"):
                pipeline.download_document(session, pipeline.discover_documents(DOCS)[0])

    def test_pdf_docx_and_spreadsheet_paths(self):
        self.assertIn("inlineData", pipeline.document_part("x.pdf", b"%PDF-test"))
        doc = Document()
        doc.add_paragraph("Delivery in 10 days")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Quantity: 5"
        stream = io.BytesIO()
        doc.save(stream)
        text = pipeline.document_part("x.docx", stream.getvalue())["text"]
        self.assertIn("10 days", text)
        self.assertIn("Quantity: 5", text)
        book = Workbook()
        book.active.append(["Quantity", 10])
        stream = io.BytesIO()
        book.save(stream)
        text = pipeline.document_part("x.xlsx", stream.getvalue())["text"]
        self.assertIn("row 1", text)
        self.assertIn("Quantity | 10", text)
        with patch.object(pipeline, "MAX_TEXT", 5):
            with self.assertRaisesRegex(pipeline.ProcessingError, "text limit"):
                pipeline.document_part("x.xlsx", stream.getvalue())

    def test_analysis_records_real_usage_and_includes_thinking_in_total(self):
        with patch.object(self.ai, "post", side_effect=[
            response({"totalTokens": 50}), generation(),
        ]):
            self.assertEqual(self.ai.analyze("a.pdf", {"text": "requirements"}), ANALYSIS)
        bucket = self.ai.bucket()
        self.assertEqual((bucket["input"], bucket["output"], bucket["total"]), (50, 20, 80))
        self.assertEqual(bucket["unknown_usage"], 0)
        self.assertEqual(pipeline.load_state(self.path)["usage"], self.state["usage"])

    def test_analysis_schema_requests_only_concise_table_fields(self):
        self.assertEqual(pipeline.SCHEMA["required"], list(pipeline.TABLE_FIELDS))
        self.assertEqual(set(pipeline.SCHEMA["properties"]), set(pipeline.TABLE_FIELDS))
        self.assertIn("80 characters", pipeline.PROMPT)
        self.assertIn("Return only those fields", pipeline.PROMPT)
        self.assertEqual(
            pipeline.validate_analysis({**ANALYSIS, "unused": "extra"}),
            ANALYSIS,
        )

    def test_daily_and_per_run_budgets_prevent_calls(self):
        self.ai.request_limit = 2
        self.ai.bucket()["requests"] = 2
        with patch.object(self.ai, "post") as post:
            with self.assertRaisesRegex(pipeline.Deferred, "daily request"):
                self.ai.analyze("a.pdf", {})
            post.assert_not_called()
        self.ai.request_limit = 0
        self.ai.calls = self.ai.max_requests
        with self.assertRaisesRegex(pipeline.Deferred, "Per-run"):
            self.ai.analyze("a.pdf", {})

    def test_token_budget_and_oversized_minute_limit(self):
        self.ai.token_limit = 40
        with patch.object(self.ai, "post", return_value=response({"totalTokens": 50})) as post:
            with self.assertRaisesRegex(pipeline.Deferred, "input-token budget"):
                self.ai.analyze("a.pdf", {})
            self.assertEqual(post.call_count, 1)
        self.ai.token_limit = 0
        self.ai.minute_tokens = 40
        with patch.object(self.ai, "post", return_value=response({"totalTokens": 50})):
            with self.assertRaisesRegex(pipeline.Deferred, "tokens/minute"):
                self.ai.analyze("a.pdf", {})

    def test_429_stops_further_calls_and_is_reported(self):
        with patch.object(self.ai, "post", return_value=response(status=429)) as post:
            for _ in range(2):
                with self.assertRaisesRegex(pipeline.Deferred, "429"):
                    self.ai.analyze("a.pdf", {})
            self.assertEqual(post.call_count, 1)
        self.assertEqual(self.ai.bucket()["quota_errors"], 1)
        plain, _, _ = pipeline.report(self.state, self.ai, [])
        self.assertIn("დღეს დაფიქსირებული სიხშირის/კვოტის შეცდომები (429): 1", plain)

    def test_network_failure_keeps_attempt_and_unknown_usage(self):
        with patch.object(self.ai, "post", side_effect=[
            response({"totalTokens": 50}), pipeline.ProcessingError("network failure"),
        ]):
            with self.assertRaises(pipeline.ProcessingError):
                self.ai.analyze("a.pdf", {})
        self.assertEqual(self.ai.bucket()["unknown_usage"], 1)
        self.assertEqual(self.ai.bucket()["requests"], 1)

    def test_temporary_overload_is_retried_before_giving_up(self):
        overloaded = response({"error": {"status": "UNAVAILABLE"}}, status=503)
        with patch.object(self.ai.http, "post", side_effect=[
            overloaded, overloaded, response({"totalTokens": 5}),
        ]) as http_post, patch.object(pipeline.time, "sleep") as sleep:
            self.assertEqual(self.ai.post("countTokens", {}).status_code, 200)
        self.assertEqual(http_post.call_count, 3)
        self.assertEqual([c.args[0] for c in sleep.call_args_list], [10, 30])
        with patch.object(self.ai.http, "post", return_value=overloaded) as http_post, \
                patch.object(pipeline.time, "sleep"):
            self.assertEqual(self.ai.post("countTokens", {}).status_code, 503)
        self.assertEqual(http_post.call_count, len(pipeline.RETRY_DELAYS) + 1)
        with patch.object(self.ai.http, "post", side_effect=[
            requests.ConnectionError("secret-url"), response({"totalTokens": 5}),
        ]), patch.object(pipeline.time, "sleep"):
            self.assertEqual(self.ai.post("countTokens", {}).status_code, 200)
        with patch.object(self.ai.http, "post", side_effect=requests.ConnectionError("secret-url")), \
                patch.object(pipeline.time, "sleep"):
            with self.assertRaisesRegex(pipeline.ProcessingError, "network failure") as caught:
                self.ai.post("countTokens", {})
        self.assertNotIn("secret-url", str(caught.exception))

    def test_service_error_reports_sanitized_reason_and_stops_run(self):
        reply = response({"error": {
            "status": "UNAVAILABLE",
            "message": f"Temporarily unavailable\ncredential={self.ai.key}",
            "details": [{"secret": "must-not-be-printed"}],
        }}, status=503)
        with patch.object(self.ai, "post", side_effect=[
            response({"totalTokens": 50}), reply,
        ]) as post:
            with self.assertRaises(pipeline.ProcessingError) as caught:
                self.ai.analyze("a.pdf", {})
            message = str(caught.exception)
            self.assertIn("UNAVAILABLE", message)
            self.assertIn("Temporarily unavailable", message)
            self.assertNotIn(self.ai.key, message)
            self.assertNotIn("must-not-be-printed", message)
            self.assertNotIn("\n", message)
            with self.assertRaises(pipeline.Deferred):
                self.ai.analyze("other.pdf", {})
            self.assertEqual(post.call_count, 2)
        self.assertEqual(self.ai.bucket()["unknown_usage"], 1)
        self.assertEqual(self.ai.bucket()["quota_errors"], 0)

    def test_non_json_service_error_does_not_dump_response_body(self):
        reply = response(status=503, text="<html>private proxy details</html>")
        reply.json.side_effect = ValueError("Not JSON")
        with self.assertRaises(pipeline.ProcessingError) as caught:
            self.ai.check_response(reply, self.ai.bucket())
        self.assertNotIn("private proxy", str(caught.exception))
        self.assertIn("HTTP 503", str(caught.exception))

    def test_success_without_usage_metadata_is_explicitly_unknown(self):
        reply = generation()
        del reply.json.return_value["usageMetadata"]
        with patch.object(self.ai, "post", side_effect=[
            response({"totalTokens": 50}), reply,
        ]):
            self.ai.analyze("a.pdf", {})
        self.assertEqual(self.ai.bucket()["unknown_usage"], 1)
        self.assertEqual(self.ai.bucket()["total"], 0)

    def test_minute_limits_space_generation_calls(self):
        self.ai.minute_requests = 5
        self.state["last_request_at"] = 100
        with patch.object(pipeline.time, "time", return_value=110), \
                patch.object(pipeline.time, "sleep") as sleep, \
                patch.object(self.ai, "post", side_effect=[
                    response({"totalTokens": 50}), generation(),
                ]):
            self.ai.analyze("a.pdf", {})
        sleep.assert_called_once()
        self.assertAlmostEqual(sleep.call_args.args[0], 3)

    def test_truncated_or_invalid_analysis_is_not_completed(self):
        with patch.object(self.ai, "post", side_effect=[
            response({"totalTokens": 50}), generation(finish="MAX_TOKENS"),
        ]):
            with self.assertRaisesRegex(pipeline.ProcessingError, "truncated"):
                self.ai.analyze("a.pdf", {})
        with self.assertRaises(pipeline.ProcessingError):
            pipeline.validate_analysis({"summary": "x", "findings": [{}], "limitations": []})

    def test_daily_usage_rolls_over_at_pacific_midnight(self):
        from datetime import datetime

        with patch.object(pipeline, "datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 22, 23, 59, tzinfo=pipeline.PACIFIC)
            self.ai.bucket()["requests"] = 5
            clock.now.return_value = datetime(2026, 10, 23, 0, 0, tzinfo=pipeline.PACIFIC)
            self.assertEqual(self.ai.bucket()["requests"], 0)
        self.assertEqual(len(self.state["usage"]), 2)

    def test_meter_unknown_and_exceeded_not_false_remaining_balance(self):
        self.assertIn("ლიმიტი მითითებული არ არის", pipeline.meter(4, 0))
        self.assertIn("[#####-----]", pipeline.meter(5, 10))
        self.assertIn("120%", pipeline.meter(12, 10))
        self.assertIn("ლიმიტი მიღწეულია", pipeline.meter(12, 10))

    def test_daily_report_html_escaping_and_pending(self):
        self.state["tenders"]["123"] = {
            "details": {**TENDER, "buyer": "<script>bad</script>"},
            "documents": {"app/42": {
                "name": "spec.pdf", "url": TENDER["link"],
                "status": "pending", "error": "Quota reached",
            }},
        }
        plain, rich, notified = pipeline.report(self.state, self.ai, ["CPV search failed"], {
            "123": self.state["tenders"]["123"]["details"],
        })
        self.assertIn("დამუშავების მოლოდინში მყოფი დოკუმენტების რაოდენობა: 1", plain)
        self.assertIn("ძიების ხარვეზები", plain)
        self.assertNotIn("<script>", rich)
        self.assertIn("&lt;script&gt;", rich)
        self.assertIn("123", notified)

    def test_incremental_cache_amendments_and_baseline(self):
        entry = {"details": TENDER, "documents": {}}
        session = MagicMock()
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=DOCS if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", return_value=b"%PDF-v1") as download, \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai)
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai)
            self.assertEqual(analyze.call_count, 1)
            download.return_value = b"%PDF-v2"
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai)
            self.assertEqual(analyze.call_count, 2)
            download.return_value = b"%PDF-v3"
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai, baseline=True)
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai)
            self.assertEqual(analyze.call_count, 2)
            download.return_value = b"%PDF-v4"
            pipeline.process_tender(session, TENDER, entry, self.state, self.ai)
            self.assertEqual(analyze.call_count, 3)

    def test_new_analysis_version_refreshes_legacy_cached_results(self):
        legacy_key = f"{self.ai.model}/{pipeline.digest(b'%PDF-v1')}"
        self.state["analyses"][legacy_key] = dict(ANALYSIS)
        entry = {"details": TENDER, "documents": {}}
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=DOCS if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", return_value=b"%PDF-v1"), \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
        analyze.assert_called_once()
        self.assertIn(pipeline.ANALYSIS_VERSION, entry["documents"]["app/42"]["analysis_key"])

    def test_all_documents_of_a_tender_use_one_gemini_request_and_shared_result(self):
        markup = DOCS + '<a href="library/files.php?mode=que&amp;file=43">answers.pdf</a>'
        entry = {"details": TENDER, "documents": {}}
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=markup if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", side_effect=[b"%PDF-a", b"%PDF-b"] * 2), \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
        analyze.assert_called_once()
        self.assertEqual(len(analyze.call_args.args[1]), 2)
        self.assertEqual(len(self.state["analyses"]), 1)
        self.assertEqual({d["status"] for d in entry["documents"].values()}, {"completed"})
        self.assertEqual(len({d["analysis_key"] for d in entry["documents"].values()}), 1)

    def test_documents_over_the_request_size_limit_are_left_for_manual_review(self):
        markup = DOCS + '<a href="library/files.php?mode=app&amp;file=43">second.pdf</a>'
        entry = {"details": TENDER, "documents": {}}
        with patch.object(pipeline, "MAX_REQUEST_BYTES", 10), \
                patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                             response(text=markup if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", side_effect=[b"%PDF-1234", b"%PDF-5678"] * 2), \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            self.assertEqual(len(analyze.call_args.args[1]), 1)
            statuses = sorted(d["status"] for d in entry["documents"].values())
            self.assertEqual(statuses, ["completed", "skipped"])
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            self.assertEqual(analyze.call_count, 1)
            self.assertEqual(sorted(d["status"] for d in entry["documents"].values()),
                             ["completed", "skipped"])

    def test_failed_combined_analysis_marks_all_documents_pending(self):
        markup = DOCS + '<a href="library/files.php?mode=app&amp;file=43">second.pdf</a>'
        entry = {"details": TENDER, "documents": {}}
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=markup if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", side_effect=[b"%PDF-a", b"%PDF-b"]), \
                patch.object(self.ai, "analyze", side_effect=pipeline.ProcessingError("Gemini HTTP 503")):
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
        self.assertEqual({d["status"] for d in entry["documents"].values()}, {"pending"})
        self.assertEqual(self.state["analyses"], {})

    def test_signed_link_changes_do_not_trigger_repeat_notifications(self):
        entry = {"details": TENDER, "documents": {"app/42": {
            "name": "x.pdf", "hash": "abc", "url": "old", "status": "completed",
        }}}
        original = pipeline.notification_fingerprint(entry)
        entry["documents"]["app/42"]["url"] = "new"
        self.assertEqual(original, pipeline.notification_fingerprint(entry))

    def test_full_daily_pipeline_analyzes_once_reports_updates_and_saves_delivery(self):
        session = MagicMock()
        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                             response(text=DOCS if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", return_value=b"%PDF-v1") as download, \
                patch.object(pipeline.Gemini, "post", side_effect=[
                    response({"totalTokens": 50}), generation(),
                    response({"totalTokens": 50}), generation(),
                ]) as post, \
                patch.object(pipeline, "send_report") as send:
            search = lambda *args: [TENDER]
            pipeline.run_monitor(session, search, {"code": "CPV"})
            self.assertIn("NAT123", send.call_args.args[0])
            self.assertIn("ახალი ან განახლებული ჩანაწერების რაოდენობა: 1", send.call_args.args[0])
            self.assertNotIn("Summary", send.call_args.args[0])
            self.assertEqual(post.call_count, 2)
            saved = pipeline.load_state(self.path)
            self.assertEqual(len(saved["analyses"]), 1)
            self.assertIn("notified", saved["tenders"]["123"])
            pipeline.run_monitor(session, search, {"code": "CPV"})
            self.assertEqual(post.call_count, 2)
            self.assertIn("ახალი ტენდერები ან დოკუმენტაციის ცვლილებები არ გამოვლენილა.", send.call_args.args[0])
            download.return_value = b"%PDF-amended"
            pipeline.run_monitor(session, search, {"code": "CPV"})
            self.assertEqual(post.call_count, 4)
            self.assertIn("ახალი ან განახლებული ჩანაწერების რაოდენობა: 1", send.call_args.args[0])
            self.assertEqual(len(pipeline.load_state(self.path)["analyses"]), 2)

    def test_failed_document_is_reported_and_retried_using_cached_other_document(self):
        entry = {"details": TENDER, "documents": {}}
        markup = DOCS + '<a href="library/files.php?mode=app&amp;file=43">other.pdf</a>'
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=markup if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", side_effect=[
                    b"%PDF-v1", requests.ConnectionError("Download failed"),
                    b"%PDF-v1", b"%PDF-v2",
                ]), \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            self.assertEqual(entry["documents"]["app/43"]["status"], "pending")
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            self.assertEqual(entry["documents"]["app/43"]["status"], "completed")
            self.assertEqual(analyze.call_count, 2)

    def test_empty_and_unrecognized_documentation_are_reported_as_incomplete(self):
        entry = {"details": TENDER, "documents": {}}
        with patch.object(pipeline, "portal_get", return_value=response(text="<h2>დოკუმენტაცია</h2>")):
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
        self.assertIn("No public attachments", entry["error"])
        with patch.object(pipeline, "portal_get", return_value=response(text="<html>Session error</html>")):
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
        self.assertIn("unrecognized HTML", entry["error"])

    def test_documentation_tab_text_analyzed_without_scripts(self):
        entry = {"details": TENDER, "documents": {}}
        markup = "<h2>დოკუმენტაცია</h2><p>" + "Requirements " * 40 + "</p><script>secretDynamic()</script>"
        with patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                          response(text=markup if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(self.ai, "analyze", return_value=ANALYSIS) as analyze:
            pipeline.process_tender(MagicMock(), TENDER, entry, self.state, self.ai)
            self.assertEqual(entry["documents"]["tab-text"]["status"], "completed")
            self.assertNotIn("secretDynamic", analyze.call_args.args[1][0]["text"])

    def test_email_failure_preserves_analysis_and_retries_notification(self):
        def process(session, tender, entry, state, ai, baseline):
            entry["documents"] = {}
            entry["error"] = ""

        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "process_tender", side_effect=process), \
                patch.object(pipeline, "send_report", side_effect=OSError("SMTP unavailable")):
            with self.assertRaises(OSError):
                pipeline.run_monitor(MagicMock(), lambda *args: [TENDER], {"x": "CPV"})
        saved = pipeline.load_state(self.path)
        self.assertNotIn("notified", saved["tenders"]["123"])
        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "process_tender", side_effect=process), \
                patch.object(pipeline, "send_report") as send:
            pipeline.run_monitor(MagicMock(), lambda *args: [TENDER], {"x": "CPV"})
            self.assertIn("ახალი ან განახლებული ჩანაწერების რაოდენობა: 1", send.call_args.args[0])
        self.assertIn("notified", pipeline.load_state(self.path)["tenders"]["123"])
        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "process_tender", side_effect=process), \
                patch.object(pipeline, "send_report") as send:
            pipeline.run_monitor(MagicMock(), lambda *args: [TENDER], {"x": "CPV"})
            self.assertIn("ახალი ტენდერები ან დოკუმენტაციის ცვლილებები არ გამოვლენილა.", send.call_args.args[0])

    def test_unexpected_tender_error_does_not_stop_run_or_email(self):
        second = {**TENDER, "id": "124", "reg_num": "NAT124"}
        calls = []

        def process(session, tender, entry, state, ai, baseline):
            calls.append(tender["id"])
            if tender["id"] == "123":
                raise KeyError("secret-detail")
            entry["documents"], entry["error"] = {}, ""

        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "process_tender", side_effect=process), \
                patch.object(pipeline, "send_report") as send:
            pipeline.run_monitor(MagicMock(), lambda *args: [TENDER, second], {"x": "CPV"})
        self.assertEqual(calls, ["123", "124"])
        self.assertNotIn("secret-detail", send.call_args.args[0])
        self.assertIn("მოულოდნელი შეცდომა", send.call_args.args[0])
        self.assertIn("NAT124", send.call_args.args[0] + send.call_args.args[1])

    def test_exhausted_time_budget_skips_remaining_downloads_and_still_emails(self):
        self.ai.deadline = 0
        with self.assertRaisesRegex(pipeline.Deferred, "time budget"):
            self.ai.analyze("a.pdf", {})
        with patch.dict(os.environ, {"MAX_RUN_MINUTES": "0"}), \
                patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "process_tender") as process, \
                patch.object(pipeline, "send_report") as send:
            pipeline.run_monitor(MagicMock(), lambda *args: [TENDER], {"x": "CPV"})
        process.assert_not_called()
        send.assert_called_once()

    def test_partial_discovery_failure_still_emails_then_fails_job(self):
        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "send_report") as send:
            with self.assertRaises(pipeline.ProcessingError):
                pipeline.run_monitor(MagicMock(), MagicMock(side_effect=requests.Timeout()), {"x": "CPV"})
        self.assertIn("ძიების ხარვეზები", send.call_args.args[0])
        self.assertIn("ძიება სრულად ვერ დასრულდა", send.call_args.args[0])
        self.assertNotIn("ახალი ტენდერები ან დოკუმენტაციის ცვლილებები არ გამოვლენილა.", send.call_args.args[0])
        self.assertNotIn("initialized", pipeline.load_state(self.path))

    def test_gemini_503_still_sends_daily_email_with_reason_and_retry_notice(self):
        with patch.object(pipeline.time, "sleep"), \
                patch.object(pipeline, "portal_get", side_effect=lambda *a, **kw:
                             response(text=DOCS if kw.get("params", {}).get("action") == "app_docs" else "Public")), \
                patch.object(pipeline, "download_document", return_value=b"%PDF-test"), \
                patch.object(pipeline.Gemini, "post", side_effect=[
                    response({"totalTokens": 50}),
                    response({"error": {
                        "status": "UNAVAILABLE",
                        "message": "This model is currently experiencing high demand.",
                    }}, status=503),
                ]), \
                patch.object(pipeline.smtplib, "SMTP_SSL") as smtp:
            server = smtp.return_value.__enter__.return_value
            server.send_message.return_value = {}
            pipeline.run_monitor(MagicMock(), lambda *args: [TENDER], {"x": "CPV"})
        server.send_message.assert_called_once()
        message = server.send_message.call_args.args[0]
        for part in message.iter_parts():
            body = part.get_content()
            self.assertIn(TENDER["link"], body)
            if part.get_content_type() == "text/html":
                body = BeautifulSoup(body, "html.parser").get_text()
            for expected in (
                "NAT123", "დროებით მიუწვდომელია",
                "დამუშავების მოლოდინში მყოფი დოკუმენტების რაოდენობა: 1",
                "შემდეგი დაგეგმილი შემოწმებისას",
            ):
                self.assertIn(expected, body)
            self.assertNotIn(ENV["GEMINI_API_KEY"], body)
        saved = pipeline.load_state(self.path)
        self.assertEqual(saved["tenders"]["123"]["documents"]["app/42"]["status"], "pending")
        self.assertIn("notified", saved["tenders"]["123"])

    def test_report_numbers_and_counts_active_tenders_separately_from_old_pending(self):
        current = {**TENDER, "id": "456", "reg_num": "NAT456",
                   "link": pipeline.PORTAL + "?go=456"}
        self.state["tenders"]["123"] = {
            "details": TENDER, "documents": {"app/42": {
                "name": "old.pdf", "url": TENDER["link"],
                "status": "pending", "error": "Per-run analysis budget reached",
            }},
        }
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"456": current})
        self.assertIn(current["link"], rich)
        self.assertIn(current["link"], plain)
        for body in (plain, BeautifulSoup(rich, "html.parser").get_text()):
            self.assertIn("ამ შემოწმებისას მოძიებული აქტიური IT ტენდერების რაოდენობა: 1", body)
            self.assertIn("NAT456", body)
            self.assertNotIn("NAT123", body)
            self.assertIn("ახალი ან განახლებული ჩანაწერების რაოდენობა: 1", body)
            self.assertIn("დამუშავების მოლოდინში მყოფი დოკუმენტების რაოდენობა: 1", body)

    def test_unchanged_tender_remains_in_overview_without_repeated_analysis(self):
        entry = {"details": TENDER, "documents": {}, "error": ""}
        entry["notified"] = pipeline.notification_fingerprint(entry)
        self.state["tenders"]["123"] = entry
        plain, _, notified = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        self.assertIn("NAT123", plain)
        self.assertIn("ახალი ან განახლებული ჩანაწერების რაოდენობა: 0", plain)
        self.assertEqual(notified, {})
        partial, _, _ = pipeline.report(self.state, self.ai, ["Connection reset"], {})
        self.assertIn("0 (არასრული შემოწმება)", partial)
        self.assertIn("ძიება სრულად ვერ დასრულდა", partial)

    def test_email_has_plain_and_html_parts_and_raises_delivery_errors(self):
        with patch.object(pipeline.smtplib, "SMTP_SSL") as smtp:
            server = smtp.return_value.__enter__.return_value
            server.send_message.return_value = {}
            pipeline.send_report("plain", "<p>html</p>")
            message = server.send_message.call_args.args[0]
            self.assertRegex(
                str(message["Subject"]),
                r"^IT ტენდერების მონიტორინგი - 72 საათიანი ანგარიში \| "
                r"\d{4}-\d{2}-\d{2} - \d{2}:\d{2}$",
            )
            self.assertEqual([p.get_content_type() for p in message.iter_parts()],
                             ["text/plain", "text/html"])
            server.send_message.return_value = {"recipient@example.test": (550, b"refused")}
            with self.assertRaises(pipeline.ProcessingError):
                pipeline.send_report("plain", "<p>html</p>")

    def test_georgian_report_preserves_sources_and_technical_errors(self):
        self.state["analyses"]["cached"] = {
            "summary": "დოკუმენტში განსაზღვრულია საქონლის მიწოდების პირობები.",
            "findings": [{"topic": "მიწოდება", "detail": "მიწოდების ვადა შეადგენს 10 დღეს.",
                          "source": "პუნქტი 2"}],
            "limitations": ["ფასების დანართი ხელმისაწვდომი არ არის."],
        }
        self.state["tenders"]["123"] = {
            "details": TENDER,
            "notice": "Portal mentions confidential material. Public files may be incomplete.",
            "documents": {
                "tab-text": {"id": "tab-text", "name": "Documentation tab (text)",
                             "url": TENDER["link"], "status": "completed", "analysis_key": "cached"},
                "app/42": {"id": "app/42", "name": "spec.pdf", "url": TENDER["link"],
                           "status": "pending", "error": "Per-run analysis budget reached"},
                "app/43": {"id": "app/43", "name": "baseline.pdf", "url": TENDER["link"],
                           "status": "baseline"},
            },
        }
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        for body in (plain, rich):
            for text in (
                "IT ტენდერების მონიტორინგის ანგარიში",
                "საიდუმლო მასალების არსებობა", "ამ გაშვებისთვის განსაზღვრული",
                "ტოკენების დაუზუსტებელი მოხმარებით",
            ):
                self.assertIn(text, body)
            self.assertNotIn("Documentation tab (text)", body)
            self.assertNotIn("Portal mentions confidential material", body)
            self.assertNotIn("მოკლე მიმოხილვა:", body)
            self.assertNotIn("ტექნიკური ინფორმაცია:", body)
        self.assertIn('lang="ka"', rich)
        self.assertIn("keywords", pipeline.PROMPT)

    def test_report_replaces_em_dashes_in_cached_and_external_text(self):
        self.state["analyses"]["cached"] = {
            "summary": "Summary\u2014detail", "findings": [], "limitations": [],
            "subject": "Subject\u2014detail", "delivery_deadline": None, "warranty": None,
        }
        tender = {**TENDER, "buyer": "Buyer\u2014name"}
        self.state["tenders"]["123"] = {
            "details": tender, "documents": {"app/42": {
                "name": "file\u2014name.pdf", "url": TENDER["link"],
                "status": "completed", "analysis_key": "cached",
            }},
        }
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": tender})
        for body in (plain, rich):
            self.assertNotIn("\u2014", body)
            self.assertNotIn("Summary-detail", body)
            self.assertIn("Subject-detail", body)
            self.assertIn("Buyer-name", body)
            self.assertNotIn("file-name.pdf", body)

    def test_georgian_error_explanations_keep_original_diagnostics(self):
        cases = (
            ("Gemini HTTP 503: UNAVAILABLE | high demand", "დროებით მიუწვდომელია"),
            ("Gemini HTTP 429", "კვოტის"),
            ("Expanded office document exceeds safety limit", "უსაფრთხოების ზღვარს"),
            ("Unsupported file or HTML/session error", "ფორმატი არ არის მხარდაჭერილი"),
            ("Unknown portal error", "დამუშავება ვერ დასრულდა"),
        )
        for error, expected in cases:
            with self.subTest(error=error):
                text = pipeline.report_error(error)
                self.assertIn(expected, text)
                self.assertIn("ტექნიკური ინფორმაცია: " + error, text)

    def test_html_priorities_use_current_scan_and_tbilisi_date(self):
        from datetime import datetime

        def tender(identifier, deadline):
            return {
                **TENDER, "id": identifier, "reg_num": identifier,
                "deadline": deadline, "link": pipeline.PORTAL + "?go=" + identifier,
            }

        current = {
            "late": tender("late", "12.11.2026"),
            "medium": tender("medium", "2026-11-01"),
            "near": tender("near", "29/10/2026"),
            "today": tender("today", "22.10.2026"),
            "invalid": tender("invalid", "31.02.2026"),
            "past": tender("past", "21.10.2026"),
        }
        self.state["tenders"]["historical"] = {
            "details": tender("historical", "22.10.2026"), "documents": {},
        }
        with patch.object(pipeline, "datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 22, 0, 30)
            clock.side_effect = datetime
            _, rich, _ = pipeline.report(self.state, self.ai, [], current)
        # The tables must not contain historical tenders, even when their updates
        # are still included in the complete report below.
        soup = BeautifulSoup(rich, "html.parser")
        tables = soup.select("table.tender-table")
        self.assertEqual(len(tables), 3)
        expected = (["today", "near"], ["medium"], ["past", "late", "invalid"])
        for table, identifiers in zip(tables, expected):
            self.assertEqual(
                [row.find_all("td")[1].get_text() for row in table.select("tbody tr")],
                identifiers,
            )
            self.assertNotIn("historical", table.get_text())
        self.assertEqual(clock.now.call_args.args[0].key, "Asia/Tbilisi")

    def test_html_keeps_complete_plain_report_and_escapes_table_content(self):
        tender = {
            **TENDER, "buyer": '<script>alert("x")</script>',
            "link": 'javascript:alert("x")',
        }
        plain, rich, _ = pipeline.report(
            self.state, self.ai, ["Connection reset"], {"123": tender},
        )
        soup = BeautifulSoup(rich, "html.parser")
        for paragraph in soup.select("#report-summary tr, #report-summary h2"):
            self.assertIn(paragraph.get_text(), plain)
        self.assertIsNone(soup.find("pre"))
        self.assertIsNone(soup.find("script"))
        self.assertEqual(soup.find_all("a"), [])
        self.assertIn("არასრული შემოწმება", plain)

    def test_all_table_columns_and_structured_values_in_both_formats(self):
        result = {
            **ANALYSIS, "subject": "კომპიუტერები (პუნქტი 1)",
            "delivery_deadline": "30 დღე (პუნქტი 2)",
            "warranty": "24 თვე (პუნქტი 3)",
        }
        self.state["analyses"]["new"] = result
        self.state["tenders"]["123"] = {
            "details": TENDER, "notice": "Confidential notice",
            "documents": {"x": {
                "name": "spec.pdf", "status": "completed", "analysis_key": "new",
            }},
        }
        # Mark unchanged to exercise overview rendering without document updates.
        entry = self.state["tenders"]["123"]
        entry["notified"] = pipeline.notification_fingerprint(entry)
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        table = BeautifulSoup(rich, "html.parser").find("table")
        headers = [th.get_text() for th in table.find_all("th")]
        self.assertEqual(len(headers), 11)
        for label in (
            "შესყიდვის საგანი", "მიწოდების ვადა", "საგარანტიო პირობები",
            "საიდუმლო მასალების შესახებ მითითება",
        ):
            self.assertIn(label, headers)
            self.assertIn(label, plain)
        for value in result["subject"], result["delivery_deadline"], result["warranty"]:
            self.assertIn(value, plain)
            self.assertIn(value, table.get_text())
        self.assertNotIn("spec.pdf", table.get_text())
        self.assertIn("დიახ", table.get_text())
        self.assertEqual(len(table.select_one("tbody tr").find_all("td")), 11)

    def test_compact_report_omits_both_removed_sections_and_groups_errors(self):
        docs = {
            str(i): {"name": f"private-detail-{i}.pdf", "status": "pending",
                     "error": "Gemini HTTP 503: UNAVAILABLE | high demand"}
            for i in range(100)
        }
        self.state["tenders"]["123"] = {"details": TENDER, "documents": docs}
        plain, rich, notified = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        for body in (plain, rich):
            self.assertNotIn("#1 | NAT123", body)
            self.assertNotIn("ტენდერების განახლებები და დასამუშავებელი დოკუმენტები", body)
            self.assertNotIn("private-detail-", body)
            self.assertNotIn("UNAVAILABLE", body)
            self.assertEqual(body.count("დროებით მიუწვდომელია"), 1)
            self.assertIn("შემთხვევების რაოდენობა: 100", body)
        for body in (plain, BeautifulSoup(rich, "html.parser").get_text()):
            self.assertIn("დოკუმენტების რაოდენობა: 100", body)
        soup = BeautifulSoup(rich, "html.parser")
        self.assertEqual(len(soup.select("table.tender-table tbody tr")), 1)
        self.assertEqual(len(soup.select("table.tender-table th")), 11)
        self.assertNotIn("NAT123", soup.select_one("#report-summary").get_text())
        self.assertIn("123", notified)
        self.assertEqual(self.state["tenders"]["123"]["documents"], docs)

    def test_table_fields_distinguish_legacy_pending_absent_and_conflicting_sources(self):
        self.state["analyses"].update({
            "legacy": {"summary": "Old", "findings": [], "limitations": []},
            "missing": {**ANALYSIS},
            "a": {**ANALYSIS, "warranty": "12 თვე (პუნქტი 1)"},
            "b": {**ANALYSIS, "warranty": "24 თვე (პუნქტი 2)"},
        })
        def entry(*keys):
            return {"documents": {
                key: {"status": "completed", "analysis_key": key, "name": key + ".pdf"}
                for key in keys
            }}
        fields = pipeline.table_analysis_fields(entry("legacy"), self.state["analyses"])
        self.assertIn("ძველ ანალიზში", fields["subject"])
        self.assertIn("არარსებობა დაუდასტურებელია", fields["confidential"])
        fields = pipeline.table_analysis_fields({}, self.state["analyses"])
        self.assertIn("ანალიზი დაუსრულებელია", fields["subject"])
        fields = pipeline.table_analysis_fields(entry("missing"), self.state["analyses"])
        self.assertIn("დოკუმენტში მითითებული არ არის", fields["subject"])
        fields = pipeline.table_analysis_fields(entry("a", "b"), self.state["analyses"])
        self.assertIn("12 თვე", fields["warranty"])
        self.assertIn("24 თვე", fields["warranty"])
        self.assertIn("წყაროებში განსხვავებაა", fields["warranty"])

    def test_skipped_file_does_not_make_tender_incomplete(self):
        self.state["analyses"]["a"] = {**ANALYSIS, "warranty": "12 თვე"}
        entry = {"documents": {
            "app/1": {"status": "completed", "analysis_key": "a"},
            "que/2": {"status": "skipped", "error": "Expanded office document exceeds safety limit"},
        }}
        fields = pipeline.table_analysis_fields(entry, self.state["analyses"])
        self.assertEqual(fields["warranty"], "12 თვე")
        self.assertNotIn("დაუსრულებელია", "".join(fields.values()))
        self.assertIn("ხელით შესამოწმებელია", fields["subject"])

    def test_gmail_layout_has_inline_tables_readable_summary_and_usage_bars(self):
        self.ai.request_limit = 20
        self.ai.bucket()["requests"] = 4
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        soup = BeautifulSoup(rich, "html.parser")
        table = soup.find("table")
        self.assertEqual(table["border"], "1")
        self.assertEqual(len(table.find_all("th", scope="col")), 11)
        self.assertNotIn("overflow-x", table.parent["style"])
        self.assertIsNone(soup.find("pre"))
        self.assertIsNone(soup.find("script"))
        self.assertIsNone(soup.find("link"))
        summary = soup.select_one("#report-summary")
        self.assertGreater(len(summary.find_all("tr")), 5)
        bars = [td for td in summary.find_all("td") if "[##--------]" in td.get_text()]
        self.assertEqual(len(bars), 1)
        self.assertIn("monospace", bars[0]["style"])
        self.assertIn("4/20", plain)

    def test_summary_groups_metrics_status_and_notes_without_losing_text(self):
        self.state["tenders"]["123"] = {
            "details": TENDER, "documents": {
                "x": {"status": "pending", "error": "Gemini HTTP 503: UNAVAILABLE"},
            },
        }
        self.ai.model = "<model&test>"
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        soup = BeautifulSoup(rich, "html.parser")
        summary = soup.select_one("#report-summary")
        blocks = summary.select(".summary-section")
        self.assertEqual([b.h3.get_text() for b in blocks], [
            "შემოწმების შედეგები", "დამუშავების სტატუსი",
            "დამუშავების ხარვეზების მოკლე შეჯამება",
            "GEMINI-ის გამოყენების სტატისტიკა - მხოლოდ ამ მონიტორის მონაცემები",
            "მიმდინარე თვის გამოყენება",
            "მნიშვნელოვანი შენიშვნები",
        ])
        self.assertIn("დოკუმენტების რაოდენობა: 1", blocks[1].get_text())
        self.assertIn("შემთხვევების რაოდენობა: 1", blocks[2].get_text())
        self.assertIn("#fff7ed", blocks[2]["style"])
        self.assertIn("<model&test>", blocks[3].get_text())
        self.assertIsNone(summary.find("model"))
        self.assertEqual(len(blocks[5].select("td")), 1)
        self.assertIsNotNone(blocks[3].find("strong", string="მოდელი: "))
        table_heading = soup.find("table").find_previous("h3").get_text()
        for line in plain.partition("\n\n" + table_heading)[0].splitlines():
            if line:
                self.assertIn(line, summary.get_text())
        for node in summary.select("tr"):
            self.assertIn(node.get_text(), plain)

    def test_combined_summary_table_matches_tender_width_without_plain_text_duplication(self):
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        soup = BeautifulSoup(rich, "html.parser")
        summary_tables = soup.select("#report-summary table")
        self.assertEqual(len(summary_tables), 1)
        summary = summary_tables[0]
        tender = soup.select_one("table.tender-table")
        for table in (summary, tender):
            self.assertEqual(table["width"], "100%")
            self.assertIn("width:100%;table-layout:fixed", table["style"])
            self.assertNotIn("min-width", table["style"])
            self.assertNotIn("overflow-x", table.parent["style"])
        self.assertNotIn("max-width", soup.select_one("#report-summary")["style"])
        for row in summary.select("tr"):
            self.assertEqual(sum(int(cell.get("colspan", 1)) for cell in row.find_all(
                ["th", "td"], recursive=False,
            )), 2)
        self.assertEqual(plain.count("მოდელი: test-model"), 1)
        self.assertEqual(len(soup.select("table.tender-table th")), 11)

    def test_report_omits_quota_explanations_in_both_formats(self):
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
        for body in (plain, rich):
            for removed in (
                "კვოტისა და აღრიცხვის განმარტებები",
                "ჯამი შესაძლოა მოიცავდეს მსჯელობისა და ქეშირებულ ტოკენებს",
                "ლიმიტები მითითებულია ხელით.",
                "countTokens-ის წინასწარი დათვლის მოთხოვნები",
                "TPD აქ აღნიშნავს",
            ):
                self.assertNotIn(removed, body)
            self.assertIn("მაჩვენებლები არ ასახავს მთელი პროექტის დარჩენილ კვოტას.", body)

    def test_all_eleven_columns_receive_the_correct_source_values(self):
        source = BeautifulSoup("""
            <table><tr id="A123"><td>
            <p class="status">გამოცხადებულია</p>
            <p>განცხადების ნომერი: <strong>NAT123</strong></p>
            <p>შემსყიდველი: <strong>სატესტო ორგანიზაცია</strong></p>
            <p>შესყიდვის კატეგორია: 30200000 - კომპიუტერები</p>
            <p>სავარაუდო ღირებულება: 12 000 GEL</p>
            <p>მიღების ვადა: 30.10.2026 12:00</p>
            </td></tr></table>""", "html.parser")
        tender = monitor.parse_search_results(source, "18924")[0]
        self.state["analyses"]["result"] = {
            **ANALYSIS, "subject": "10 კომპიუტერი (პუნქტი 1)",
            "delivery_deadline": "30 დღე (პუნქტი 2)",
            "warranty": "24 თვე (პუნქტი 3)",
        }
        self.state["tenders"]["123"] = {
            "details": tender, "notice": "Confidential notice",
            "documents": {"doc": {
                "name": "spec.pdf", "status": "completed", "analysis_key": "result",
            }},
        }
        plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": tender})
        row = BeautifulSoup(rich, "html.parser").select_one("tbody tr")
        cells = row.find_all("td")
        expected = [
            "1", "NAT123", "სატესტო ორგანიზაცია", "30200000 - კომპიუტერები",
            "10 კომპიუტერი (პუნქტი 1)", "სავარაუდო ღირებულება: 12 000 GEL",
            "მიღების ვადა: 30.10.2026 12:00", "30 დღე (პუნქტი 2)",
            "24 თვე (პუნქტი 3)", "დიახ - პორტალზე მითითებულია",
        ]
        self.assertEqual(len(cells), 11)
        for cell, value in zip(cells, expected):
            self.assertIn(value, cell.get_text())
            self.assertIn(value, plain)
        self.assertEqual(cells[10].find("a")["href"], tender["link"])

    def test_missing_portal_values_are_explicit_not_blank_or_invented(self):
        _, rich, _ = pipeline.report(self.state, self.ai, [], {
            "123": {"id": "123", "link": TENDER["link"]},
        })
        cells = BeautifulSoup(rich, "html.parser").select_one("tbody tr").find_all("td")
        for index in (1, 2, 3, 5, 6):
            self.assertEqual(cells[index].get_text(), "პორტალის მონაცემებში არ გამოვლენილა.")
        self.assertIn("ანალიზი დაუსრულებელია", cells[4].get_text())

    def test_table_values_omit_document_labels_and_deduplicate_equal_values(self):
        self.state["analyses"]["same"] = {**ANALYSIS, "subject": "10 კომპიუტერი (პუნქტი 1)"}
        entry = {"documents": {
            str(i): {"name": f"Annex_{i}.pdf", "status": "completed", "analysis_key": "same"}
            for i in range(2)
        }}
        original = json.dumps(entry, sort_keys=True)
        fields = pipeline.table_analysis_fields(entry, self.state["analyses"])
        self.assertEqual(fields["subject"], "10 კომპიუტერი (პუნქტი 1)")
        self.assertNotIn("დოკუმენტი:", fields["subject"])
        self.assertEqual(json.dumps(entry, sort_keys=True), original)

    def test_report_heading_uses_tbilisi_date_and_time(self):
        from datetime import datetime, timezone
        instant = datetime(2026, 10, 22, 22, 24, tzinfo=timezone.utc)
        with patch.object(pipeline, "datetime", wraps=datetime) as clock:
            clock.now.side_effect = lambda tz: instant.astimezone(tz)
            plain, rich, _ = pipeline.report(self.state, self.ai, [], {})
        expected = (
            "IT ტენდერების მონიტორინგის ანგარიში - "
            "2026-10-23 02:24 (თბილისის დრო, UTC+04:00)"
        )
        self.assertTrue(plain.startswith(expected))
        self.assertIn(expected, BeautifulSoup(rich, "html.parser").get_text())
        self.assertIn("2026-10-22 (America/Los_Angeles)", plain)

    def test_wide_table_wraps_long_unbroken_values_without_truncation(self):
        long_value = "IPMPLS_TRANSPORT_NETWORK_" * 20
        _, rich, _ = pipeline.report(
            self.state, self.ai, [], {"123": {**TENDER, "buyer": long_value}},
        )
        table = BeautifulSoup(rich, "html.parser").find("table")
        self.assertIn("width:100%", table["style"])
        self.assertNotIn("min-width", table["style"])
        self.assertEqual(sum(int(col["width"].rstrip("%")) for col in table.find_all("col")), 100)
        self.assertIn(long_value, table.get_text())
        for cell in table.find_all(["th", "td"]):
            self.assertIn("overflow-wrap:anywhere", cell["style"])
            self.assertIn("word-break:break-word", cell["style"])
            self.assertNotIn("overflow:hidden", cell["style"])

    def test_confidential_document_analysis_reaches_both_email_formats(self):
        value = "საიდუმლო დანართი გაიცემა შესაბამისი დაშვების საფუძველზე (პუნქტი 5)"
        self.state["analyses"]["new"] = {**ANALYSIS, "confidential": value}
        self.state["tenders"]["123"] = {
            "details": TENDER, "documents": {
                "doc": {"name": "annex.pdf", "status": "completed", "analysis_key": "new"},
            },
        }
        for notice in ("", "Portal notice"):
            self.state["tenders"]["123"]["notice"] = notice
            plain, rich, _ = pipeline.report(self.state, self.ai, [], {"123": TENDER})
            cells = BeautifulSoup(rich, "html.parser").select_one(
                "table.tender-table tbody tr",
            ).find_all("td")
            self.assertIn(value, cells[9].get_text())
            self.assertIn(value, plain)
            self.assertNotIn("[დოკუმენტი:", cells[9].get_text())
            if notice:
                self.assertIn("პორტალზე მითითებულია", cells[9].get_text())

    def test_confidential_legacy_null_and_pending_remain_distinct(self):
        self.state["analyses"]["old"] = {k: v for k, v in ANALYSIS.items() if k != "confidential"}
        self.state["analyses"]["new"] = dict(ANALYSIS)
        for key, status, expected in (
            ("old", "completed", "ძველ ანალიზში"),
            ("new", "completed", "დოკუმენტში მითითებული არ არის"),
            ("new", "pending", "ანალიზი დაუსრულებელია"),
        ):
            fields = pipeline.table_analysis_fields(
                {"documents": {"doc": {"status": status, "analysis_key": key}}},
                self.state["analyses"],
            )
            self.assertIn(expected, fields["confidential"])
            self.assertIn("არარსებობა დაუდასტურებელია", fields["confidential"])

    def test_new_analysis_requires_nullable_table_fields(self):
        self.assertEqual(pipeline.validate_analysis(ANALYSIS), ANALYSIS)
        for key in pipeline.TABLE_FIELDS:
            self.assertIn(key, pipeline.SCHEMA["required"])
            self.assertTrue(pipeline.SCHEMA["properties"][key]["nullable"])
            invalid = dict(ANALYSIS)
            del invalid[key]
            with self.assertRaisesRegex(pipeline.ProcessingError, "table field"):
                pipeline.validate_analysis(invalid)
            invalid[key] = 12
            with self.assertRaises(pipeline.ProcessingError):
                pipeline.validate_analysis(invalid)


class SearchTests(unittest.TestCase):
    def test_search_fetches_all_pages_and_preserves_announced_filter(self):
        session = MagicMock()
        first = response(text="""
        <div class="pager">8 ჩანაწერი (გვერდი: 1/2)</div>
        <table><tr id="A1"><td><p class="status">Announced</p></td></tr></table>""")
        second = response(text="""
        <div class="pager">8 ჩანაწერი (გვერდი: 2/2)</div>
        <table><tr id="A2"><td><p class="status">Announced</p></td></tr></table>""")
        session.get.side_effect = [first, second]
        session.post.side_effect = requests.ConnectionError("Remote end closed connection")
        self.assertEqual([t["id"] for t in monitor.search_tenders(session, "18924")], ["1", "2"])
        session.post.assert_not_called()
        initial = session.get.call_args_list[0].kwargs["params"]
        self.assertEqual(initial["app_status"], "10")
        self.assertEqual(initial["app_basecode"], "18924")
        self.assertEqual(initial["action"], "search_app")
        self.assertEqual(session.get.call_args.kwargs["params"]["page"], 2)

    def test_search_supports_procurement_type_filter(self):
        session = MagicMock()
        session.get.return_value = response(text='<div class="pager">0 ჩანაწერი (გვერდი: 1/0)</div>')
        self.assertEqual(monitor.search_tenders(session, "x", "NAT")[0:0], [])
        self.assertEqual(session.get.call_args.kwargs["params"]["app_type"], "9")
        self.assertEqual(monitor.search_tenders(session, "x", "GEO")[0:0], [])
        self.assertEqual(session.get.call_args.kwargs["params"]["app_type"], "3")

    def test_empty_search_and_session_error_are_not_confused(self):
        session = MagicMock()
        session.get.return_value = response(text='<div class="pager">0 ჩანაწერი (გვერდი: 1/0)</div>')
        self.assertEqual(monitor.search_tenders(session, "x"), [])
        session.get.return_value = response(text="<html>Session error</html>")
        with self.assertRaises(pipeline.ProcessingError):
            monitor.search_tenders(session, "x")


if __name__ == "__main__":
    unittest.main()
